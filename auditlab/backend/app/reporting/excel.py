"""Export Excel des tableaux de bord (REP-08) : une feuille « Filtres » (période, entité,
criticité appliquées) puis une feuille par tableau, en-têtes en gras et filtres automatiques."""

import io
from typing import Any

from sqlalchemy.orm import Session

from app.models.universe import AuditableEntity

Sheet = tuple[str, list[str], list[list[Any]]]


def to_xlsx(filters: list[tuple[str, str]], sheets: list[Sheet]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Filtres"
    for label, value in filters:
        ws.append([label, value])
    for name, header, rows in sheets:
        sheet = wb.create_sheet(name[:31])
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        if rows:
            sheet.auto_filter.ref = sheet.dimensions
        for i, column in enumerate(header, start=1):
            width = max(
                [len(str(column))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None]
            )
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = min(
                max(width + 2, 10), 60
            )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def filter_rows(db: Session, described: dict[str, Any]) -> list[tuple[str, str]]:
    entity = db.get(AuditableEntity, described["entity_id"]) if described.get("entity_id") else None
    return [
        ("Du", described.get("debut") or "—"),
        ("Au", described.get("fin") or "—"),
        ("Entité", entity.nom if entity else "Toutes"),
        ("Criticité", described.get("criticite") or "Toutes"),
    ]


def direction_sheets(data: dict[str, Any]) -> list[Sheet]:
    plan = data["plan"] or {}
    return [
        (
            "Synthèse",
            ["Indicateur", "Valeur"],
            [
                ["Plan annuel", f"{plan.get('annee', '—')} v{plan.get('version', '—')}"],
                ["Missions prévues au plan", plan.get("missions_prevues")],
                ["Missions lancées", plan.get("missions_lancees")],
                ["Missions clôturées", plan.get("missions_cloturees")],
                ["Taux de réalisation du plan (%)", plan.get("taux_realisation")],
                ["Budget (jours)", data["budget_jours"]],
                ["Jours consommés", data["jours_consommes"]],
                ["Taux de consommation (%)", data["taux_consommation"]],
            ],
        ),
        (
            "Missions par statut",
            ["Statut", "Missions"],
            [[k, v] for k, v in data["missions_par_statut"].items()],
        ),
        (
            "Missions",
            [
                "Référence",
                "Mission",
                "Entité",
                "Statut",
                "Budget (j)",
                "Consommés (j)",
                "Écart (j)",
            ],
            [
                [
                    m["reference"],
                    m["titre"],
                    m["entite"],
                    m["statut"],
                    m["budget_jours"],
                    m["jours_consommes"],
                    m["ecart_jours"],
                ]
                for m in data["missions"]
            ],
        ),
    ]


def recommendations_sheets(data: dict[str, Any]) -> list[Sheet]:
    statuses = list(next(iter(data["plans_par_criticite_et_statut"].values())).keys())
    return [
        (
            "Synthèse",
            ["Indicateur", "Valeur"],
            [
                ["Plans d'action", data["plans"]],
                ["Taux de mise en œuvre (%)", data["taux_mise_en_oeuvre"]],
                ["Taux de mise en œuvre des plans échus (%)", data["taux_mise_en_oeuvre_echus"]],
                ["Plans en retard", len(data["plans_en_retard"])],
            ],
        ),
        (
            "Criticité et statut",
            ["Criticité", "Constats", *statuses],
            [
                [c, data["constats_par_criticite"][c], *[row[s] for s in statuses]]
                for c, row in data["plans_par_criticite_et_statut"].items()
            ],
        ),
        (
            "Retards par direction",
            ["Direction", "Plans en retard", "Plus ancienne échéance"],
            [
                [r["direction"], r["plans_en_retard"], r["plus_ancienne_echeance"]]
                for r in data["retards_par_direction"]
            ],
        ),
        (
            "Plans en retard",
            [
                "Mission",
                "Constat",
                "Criticité",
                "Direction",
                "Plan d'action",
                "Responsable",
                "Échéance",
                "Jours de retard",
            ],
            [
                [
                    p["mission"],
                    p["constat"],
                    p["criticite"],
                    p["direction"],
                    p["description"],
                    p["responsable"],
                    p["echeance"],
                    p["jours_de_retard"],
                ]
                for p in data["plans_en_retard"]
            ],
        ),
    ]


def mission_sheets(data: dict[str, Any]) -> list[Sheet]:
    t = data["temps"]
    return [
        (
            "Synthèse",
            ["Indicateur", "Valeur"],
            [
                ["Papiers de travail", data["papiers"]["total"]],
                ["Taux d'approbation (%)", data["papiers"]["taux_approbation"]],
                ["Notes de revue ouvertes", data["notes_ouvertes"]],
                ["Budget (jours)", t["budget_jours"]],
                ["Jours consommés", t["jours_consommes"]],
                ["Taux de consommation (%)", t["taux_consommation"]],
            ],
        ),
        (
            "Papiers de travail",
            ["Référence", "Titre", "État", "Notes ouvertes"],
            [
                [w["reference"], w["titre"], w["etat"], w["notes_ouvertes"]]
                for w in data["papiers"]["liste"]
            ],
        ),
        (
            "Temps par personne",
            ["Personne", "Jours"],
            [[p["nom"], p["jours"]] for p in t["par_personne"]],
        ),
        (
            "Temps par activité",
            ["Activité", "Jours"],
            [[a["activite"], a["jours"]] for a in t["par_activite"]],
        ),
    ]
