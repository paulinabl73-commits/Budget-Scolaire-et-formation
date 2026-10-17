"""Export des tables en CSV et en Excel (DAT-32).

CSV : formats français (séparateur « ; », virgule décimale, dates jj/mm/aaaa), UTF-8 avec BOM
pour une ouverture directe dans Excel. Excel : une feuille de données et une feuille
« Informations » (table, empreinte, filtres, date et auteur de l'export) pour la traçabilité.
"""

import datetime as dt
import tempfile
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.analytics.grid import where_clause
from app.analytics.relation import Relation
from app.analytics.store import connect, quote_ident

EXCEL_MAX_ROWS = 1_048_575  # limite d'Excel, hors ligne d'en-tête


class ExportError(ValueError):
    pass


def _select(
    rel: Relation, columns: list[str] | None, filters: dict, search: str | None, sort: list
) -> tuple[str, list[Any], list[str]]:
    names = [c["nom"] for c in rel.schema]
    selected = columns or names
    unknown = [c for c in selected + [s["colonne"] for s in sort] if c not in names]
    if unknown:
        raise ExportError(f"Colonne inconnue « {unknown[0]} »")
    where, params = where_clause(rel.schema, filters, search)
    order = ", ".join(
        f"{quote_ident(s['colonne'])} {'DESC' if s.get('sens') == 'desc' else 'ASC'} NULLS LAST"
        for s in sort
    )
    cols = ", ".join(quote_ident(c) for c in selected)
    sql = f"SELECT {cols} FROM ({rel.sql}){where}"  # noqa: S608 - noms vérifiés, valeurs en paramètres
    if order:
        sql += f" ORDER BY {order}"
    return sql, rel.params + params, selected


def to_csv(rel: Relation, columns, filters, search, sort) -> tuple[Path, int]:  # noqa: ANN001
    sql, params, selected = _select(rel, columns, filters, search, sort)
    tmp = Path(tempfile.mkdtemp(prefix="export-")) / "export.csv"
    body = tmp.with_suffix(".tmp")
    # Virgule décimale : DuckDB ne l'écrit pas lui-même, les nombres passent en texte.
    out_cols = ", ".join(
        f"replace(CAST({quote_ident(c)} AS VARCHAR), '.', ',') AS {quote_ident(c)}"
        if rel.types[c] == "decimal"
        else quote_ident(c)
        for c in selected
    )
    with connect() as con:
        con.execute(f"CREATE TEMP TABLE sortie AS {sql}", params)
        count = con.execute("SELECT count(*) FROM sortie").fetchone()[0]
        con.execute(
            f"COPY (SELECT {out_cols} FROM sortie) TO '{body}' (FORMAT CSV, HEADER, "  # noqa: S608
            "DELIMITER ';', DATEFORMAT '%d/%m/%Y', TIMESTAMPFORMAT '%d/%m/%Y %H:%M:%S')"
        )
    with tmp.open("wb") as out, body.open("rb") as src:
        out.write(b"\xef\xbb\xbf")
        while chunk := src.read(1024 * 1024):
            out.write(chunk)
    body.unlink()
    return tmp, count


def to_excel(
    rel: Relation,
    columns,
    filters,
    search,
    sort,
    info: dict[str, Any],  # noqa: ANN001
) -> tuple[Path, int]:
    sql, params, selected = _select(rel, columns, filters, search, sort)
    types = rel.types
    with connect() as con:
        count = con.execute(f"SELECT count(*) FROM ({sql})", params).fetchone()[0]  # noqa: S608
        if count > EXCEL_MAX_ROWS:
            n = f"{count:,}".replace(",", " ")
            raise ExportError(f"{n} lignes : au-delà de la limite d'Excel, exportez en CSV")
        df = con.execute(sql, params).pl()
    wb = Workbook(write_only=True)
    sheet = wb.create_sheet("Données")
    sheet.append(selected)
    formats = {}
    for i, name in enumerate(selected):
        kind = types[name]
        formats[i] = {"decimal": "#,##0.00", "entier": "#,##0", "date": "DD/MM/YYYY",
                      "date_heure": "DD/MM/YYYY HH:MM:SS"}.get(kind)  # fmt: skip
    from openpyxl.cell import WriteOnlyCell

    for row in df.iter_rows():
        cells = []
        for i, value in enumerate(row):
            if isinstance(value, Decimal):
                value = float(value) if abs(value) < Decimal("1e15") else str(value)
            cell = WriteOnlyCell(sheet, value=value)
            if formats[i] and value is not None:
                cell.number_format = formats[i]
            cells.append(cell)
        sheet.append(cells)
    meta = wb.create_sheet("Informations")
    bold = Font(bold=True)
    for key, value in info.items():
        label = WriteOnlyCell(meta, value=key)
        label.font = bold
        meta.append([label, value if isinstance(value, (str, int, float)) else str(value)])
    meta.append(["Lignes exportées", count])
    tmp = Path(tempfile.mkdtemp(prefix="export-")) / "export.xlsx"
    wb.save(tmp)
    return tmp, count


def export_info(ds, user, filters, search) -> dict[str, Any]:  # noqa: ANN001
    return {
        "Table": ds.nom,
        "Empreinte SHA-256 de la table": ds.parquet_sha256 or "(vue calculée)",
        "Lignes de la table": ds.nb_lignes,
        "Filtres": str(filters) if filters else "aucun",
        "Recherche": search or "aucune",
        "Exporté par": user.nom,
        "Exporté le": dt.datetime.now(dt.UTC).strftime("%d/%m/%Y %H:%M:%S UTC"),
    }
