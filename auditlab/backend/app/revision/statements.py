"""États financiers révisés (REV-23) et contrôles de cohérence (REV-24).

Chaque poste est présenté avant ajustements, avec l'effet des écritures acceptées, après
ajustements et pour N-1. Les règles de contrôle du référentiel (bilan équilibré, résultat
concordant, trésorerie du tableau des flux égale à celle du bilan…) sont évaluées sur les
montants révisés ; une règle peut être limitée à certains exercices (« periodes »).
"""

import io
from decimal import Decimal
from typing import Any

from sqlalchemy.orm import Session

from app.models.revision import AccountingFramework, AuditFile, StatementLine
from app.revision.mapping import file_values, postes

ZERO = Decimal(0)
ETATS = {
    "bilan_actif": "Bilan – actif",
    "bilan_passif": "Bilan – passif",
    "compte_resultat": "Compte de résultat",
    "flux_tresorerie": "Tableau des flux de trésorerie",
    "notes": "Notes annexes",
    "autre": "Autres états",
}


def assignable(line: StatementLine) -> bool:
    """Poste de détail (pas un total calculé) : il peut recevoir des comptes."""
    return bool(line.racines) or not line.formule


def checks(
    rules: list[dict[str, Any]], values: dict[str, dict[str, Decimal]]
) -> list[dict[str, Any]]:
    out = []
    for rule in rules:
        periods = rule.get("periodes") or ["N", "N-1"]
        gaps = {
            p: values[p].get(rule["gauche"], ZERO) - values[p].get(rule["droite"], ZERO)
            for p in periods
        }
        out.append(
            {
                "code": rule["code"],
                "libelle": rule["libelle"],
                "ecarts": {p: str(v) for p, v in gaps.items()},
                "ok": all(v == 0 for v in gaps.values()),
            }
        )
    return out


def build(db: Session, file: AuditFile) -> dict[str, Any]:
    """États avant / après ajustements ; UnmappedError si un compte n'est pas affecté."""
    before = file_values(db, file)
    after = file_values(db, file, adjusted=True)
    fw = db.get(AccountingFramework, file.framework_id)
    states: dict[str, list[dict[str, Any]]] = {}
    for ln in postes(db, file.framework_id):
        n_before, n_after = before["N"].get(ln.code, ZERO), after["N"].get(ln.code, ZERO)
        n1 = None if ln.etat == "flux_tresorerie" else str(after["N-1"].get(ln.code, ZERO))
        states.setdefault(ln.etat, []).append(
            {
                "code": ln.code,
                "libelle": ln.libelle,
                "niveau": ln.niveau,
                "total": not assignable(ln),
                "N_avant": str(n_before),
                "ajustements": str(n_after - n_before),
                "N": str(n_after),
                "N-1": n1,
            }
        )
    return {
        "referentiel": f"{fw.libelle} ({fw.code} {fw.version})",
        "etats": {k: states[k] for k in ETATS if k in states},
        "controles": checks(fw.regles, after),
        "controles_avant": checks(fw.regles, before),
    }


# --- Exports ------------------------------------------------------------------------------------


def _number(value: str | None) -> float | None:
    return None if value is None else float(Decimal(value))


def to_excel(data: dict[str, Any], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    headers = ["Code", "Poste", "N avant ajustements", "Ajustements", "N révisé", "N-1"]
    for etat, lines in data["etats"].items():
        ws = wb.create_sheet(ETATS[etat][:31])
        ws.append([title])
        ws.append([data["referentiel"]])
        ws.append([])
        ws.append(headers)
        for cell in ws[4]:
            cell.font = Font(bold=True)
        for ln in lines:
            ws.append(
                [
                    ln["code"],
                    ln["libelle"],
                    _number(ln["N_avant"]),
                    _number(ln["ajustements"]),
                    _number(ln["N"]),
                    _number(ln["N-1"]),
                ]
            )
            if ln["total"]:
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True)
        for col, width in zip("ABCDEF", (8, 60, 18, 16, 18, 18), strict=True):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=5, min_col=3, max_col=6):
            for cell in row:
                cell.number_format = "# ##0;-# ##0"
    ws = wb.create_sheet("Contrôles")
    ws.append(["Contrôle", "Résultat", "Écarts"])
    for c in data["controles"]:
        ws.append(
            [
                c["libelle"],
                "Conforme" if c["ok"] else "Écart",
                " ; ".join(f"{p} : {v}" for p, v in c["ecarts"].items()),
            ]
        )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _fr(value: str | None) -> str:
    if value is None:
        return ""
    return f"{Decimal(value):,.0f}".replace(",", " ")


def to_pdf(data: dict[str, Any], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table

    styles = getSampleStyleSheet()
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=title,
    )
    story: list[Any] = []
    for i, (etat, lines) in enumerate(data["etats"].items()):
        if i:
            story.append(PageBreak())
        story += [
            Paragraph(title, styles["Title"]),
            Paragraph(ETATS[etat], styles["Heading2"]),
            Paragraph(data["referentiel"], styles["Normal"]),
            Spacer(1, 4 * mm),
        ]
        rows = [["Code", "Poste", "N avant aj.", "Ajustements", "N révisé", "N-1"]]
        rows += [
            [
                ln["code"],
                Paragraph(ln["libelle"], styles["BodyText"]),
                _fr(ln["N_avant"]),
                _fr(ln["ajustements"]),
                _fr(ln["N"]),
                _fr(ln["N-1"]),
            ]
            for ln in lines
        ]
        table = Table(
            rows, colWidths=[12 * mm, 74 * mm, 24 * mm, 22 * mm, 24 * mm, 24 * mm], repeatRows=1
        )
        style = [
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
            ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        for row, ln in enumerate(lines, start=1):
            if ln["total"]:
                style += [
                    ("FONTNAME", (0, row), (-1, row), "Helvetica-Bold"),
                    ("BACKGROUND", (0, row), (-1, row), colors.whitesmoke),
                ]
        table.setStyle(style)
        story.append(table)
    story += [PageBreak(), Paragraph("Contrôles de cohérence", styles["Heading2"])]
    for c in data["controles"]:
        mark = (
            "Conforme"
            if c["ok"]
            else "ÉCART " + " ; ".join(f"{p} : {_fr(v)}" for p, v in c["ecarts"].items())
        )
        story.append(Paragraph(f"{c['libelle']} : {mark}", styles["Normal"]))
    doc.build(story)
    return out.getvalue()
