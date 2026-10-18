"""Aide à la liasse fiscale et statistique (DSF, REV-25).

Un modèle (par pays) liste des feuilles et leurs lignes : référence, libellé, renvoi à une
note et formule sur les codes de postes des états révisés (REV-23). Le modèle générique livré
reprend les états du Système normal SYSCOHADA, dont les références sont celles de la liasse ;
chaque pays ajoute ou adapte ses feuilles par import Excel. Les données non comptables
(effectifs, informations statistiques) restent à saisir dans la liasse.
"""

import io
from decimal import Decimal
from typing import Any

from app.revision.formula import FormulaError, evaluate, references, tokens

FEUILLES = {
    "bilan_actif": "Bilan – Actif",
    "bilan_passif": "Bilan – Passif",
    "compte_resultat": "Compte de résultat",
    "flux_tresorerie": "Tableau des flux de trésorerie",
}


class DsfError(ValueError):
    pass


def generic_model(postes: list[dict[str, Any]]) -> dict[str, Any]:
    """Modèle générique : une feuille par état, une ligne par poste (référence = code)."""
    sheets = []
    for etat, nom in FEUILLES.items():
        lines = [
            {"ref": p["code"], "libelle": p["libelle"], "formule": p["code"], "note": None}
            for p in postes
            if p["etat"] == etat
        ]
        sheets.append({"nom": nom, "lignes": lines, "n1": etat != "flux_tresorerie"})
    return {
        "identification": [
            "Dénomination",
            "Exercice",
            "Date de clôture",
            "Devise",
            "Référentiel",
            "Pays",
        ],
        "feuilles": sheets,
    }


def check(model: dict[str, Any], known: set[str]) -> None:
    sheets = model.get("feuilles")
    if not isinstance(sheets, list) or not sheets:
        raise DsfError("Modèle sans feuille")
    for sheet in sheets:
        if not sheet.get("nom") or not isinstance(sheet.get("lignes"), list):
            raise DsfError("Feuille sans nom ou sans lignes")
        for line in sheet["lignes"]:
            if not line.get("ref") or not line.get("libelle"):
                raise DsfError(f"{sheet['nom']} : ligne sans référence ou libellé")
            if line.get("formule"):
                try:
                    tokens(line["formule"])
                except FormulaError as exc:
                    raise DsfError(f"{line['ref']} : {exc}") from exc
                unknown = references(line["formule"]) - known
                if unknown:
                    raise DsfError(f"{line['ref']} : poste inconnu « {sorted(unknown)[0]} »")


def fill(
    model: dict[str, Any], n: dict[str, Decimal], n1: dict[str, Decimal]
) -> list[dict[str, Any]]:
    out = []
    for sheet in model["feuilles"]:
        lines = []
        for line in sheet["lignes"]:
            formula = line.get("formule")
            value_n = evaluate(formula, n) if formula else None
            value_n1 = evaluate(formula, n1) if formula and sheet.get("n1", True) else None
            lines.append(
                {
                    "ref": line["ref"],
                    "libelle": line["libelle"],
                    "note": line.get("note"),
                    "N": None if value_n is None else str(value_n),
                    "N-1": None if value_n1 is None else str(value_n1),
                }
            )
        out.append({"nom": sheet["nom"], "lignes": lines})
    return out


def to_excel(identification: dict[str, str], sheets: list[dict[str, Any]], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Identification"
    ws.append([title])
    ws["A1"].font = Font(bold=True)
    for key, value in identification.items():
        ws.append([key, value])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50
    for sheet in sheets:
        ws = wb.create_sheet(sheet["nom"][:31])
        ws.append(["Réf.", "Libellé", "Note", "Exercice N", "Exercice N-1"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for line in sheet["lignes"]:
            ws.append(
                [
                    line["ref"],
                    line["libelle"],
                    line["note"],
                    float(line["N"]) if line["N"] is not None else None,
                    float(line["N-1"]) if line["N-1"] is not None else None,
                ]
            )
        for col, width in zip("ABCDE", (8, 60, 8, 18, 18), strict=True):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
            for cell in row:
                cell.number_format = "# ##0;-# ##0"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def model_to_excel(model: dict[str, Any]) -> bytes:
    """Gabarit d'un modèle : une feuille par feuille de liasse (ref, libelle, note, formule)."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in model["feuilles"]:
        ws = wb.create_sheet(sheet["nom"][:31])
        ws.append(["ref", "libelle", "note", "formule"])
        for line in sheet["lignes"]:
            ws.append([line["ref"], line["libelle"], line.get("note"), line.get("formule")])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(value: Any) -> str | None:
    return None if value in (None, "") else str(value)


def model_from_excel(content: bytes) -> dict[str, Any]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - fichier illisible
        raise DsfError("Classeur Excel illisible") from exc
    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
        if not {"ref", "libelle"} <= set(header):
            raise DsfError(f"Feuille « {ws.title} » : colonnes ref et libelle attendues")
        idx = {h: i for i, h in enumerate(header)}
        lines = []
        for r in rows[1:]:
            if not r or r[idx["ref"]] in (None, ""):
                continue
            cells = {k: (r[i] if i < len(r) else None) for k, i in idx.items()}
            lines.append(
                {
                    "ref": str(cells["ref"]),
                    "libelle": str(cells.get("libelle") or ""),
                    "note": _text(cells.get("note")),
                    "formule": _text(cells.get("formule")),
                }
            )
        sheets.append({"nom": ws.title, "lignes": lines, "n1": "flux" not in ws.title.lower()})
    return {"identification": generic_model([])["identification"], "feuilles": sheets}
