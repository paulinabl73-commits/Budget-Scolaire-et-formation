"""Référentiels comptables (REV-01) : chargement contrôlé, import et export Excel.

Classeur Excel attendu (une feuille par rubrique, en-têtes en première ligne) :
- « Référentiel » : deux colonnes clé / valeur (code, version, libelle, pays, date_effet, source) ;
- « Comptes » : numero, libelle, classe, sens, branche ;
- « Postes » : etat, code, libelle, nature, racines, formule, ordre, niveau ; « racines » liste
  les racines séparées par « ; », suivies de « :debit » ou « :credit » pour un compte à double
  sens (ex. « 409;40:debit ») ;
- « Ratios » : code, libelle, formule, unite, famille ; « Bases » : code, libelle, formule,
  pct_min, pct_max (fourchette indicative du pourcentage de seuil, REV-11) ;
- « Règles » : code, libelle, gauche, droite ; « Paramètres » : code, libelle, valeur, unite,
  pays, exercice, branche ; « Branches » : code, libelle (CIPRES).
"""

import datetime as dt
import io
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.revision import (
    AccountingFramework,
    FrameworkAccount,
    RegulatoryParameter,
    StatementLine,
)
from app.revision.formula import PREVIOUS, FormulaError, references, tokens

CODES = ("SYSCOHADA", "IFRS", "CIPRES")
ETATS = ("bilan_actif", "bilan_passif", "compte_resultat", "flux_tresorerie", "notes", "autre")
NATURES = ("actif", "passif", "charge", "produit", "flux")
DATA_DIR = Path(__file__).parent / "data"


class FrameworkError(ValueError):
    pass


def _check(doc: dict[str, Any]) -> None:
    if doc.get("code") not in CODES:
        raise FrameworkError(f"Code de référentiel attendu parmi {', '.join(CODES)}")
    for key in ("version", "libelle", "date_effet"):
        if not doc.get(key):
            raise FrameworkError(f"Champ « {key} » obligatoire")
    numbers = set()
    for acc in doc.get("comptes", []):
        num = str(acc.get("numero", "")).strip()
        if not num.isdigit():
            raise FrameworkError(f"Numéro de compte invalide « {num} »")
        if num in numbers:
            raise FrameworkError(f"Compte « {num} » en double")
        numbers.add(num)
    codes = [p["code"] for p in doc.get("postes", [])]
    if len(set(codes)) != len(codes):
        raise FrameworkError("Code de poste en double")
    known = set(codes)
    for p in doc.get("postes", []):
        if p.get("etat") not in ETATS or p.get("nature") not in NATURES:
            raise FrameworkError(f"Poste {p.get('code')} : état ou nature invalide")
        for r in p.get("racines", []):
            if not str(r.get("racine", "")).isdigit() or r.get("sens") not in (
                None,
                "debit",
                "credit",
            ):
                raise FrameworkError(f"Poste {p['code']} : racine invalide {r}")
    formulas = [(p["code"], p.get("formule")) for p in doc.get("postes", [])]
    formulas += [(r["code"], r["formule"]) for r in doc.get("ratios", [])]
    formulas += [(b["code"], b["formule"]) for b in doc.get("bases_seuil", [])]
    formulas += [(r["code"], f"{r['gauche']}-{r['droite']}") for r in doc.get("regles", [])]
    for owner, formula in formulas:
        if not formula:
            continue
        try:
            tokens(formula)
        except FormulaError as exc:
            raise FrameworkError(f"{owner} : {exc}") from exc
        # « CODE_N1 » : montant N-1 du poste (tableau des flux, REV-24).
        unknown = {r.removesuffix(PREVIOUS) for r in references(formula)} - known
        if unknown:
            raise FrameworkError(
                f"{owner} : poste inconnu « {sorted(unknown)[0]} » dans la formule"
            )


def load_framework(
    db: Session, doc: dict[str, Any], sha256: str | None = None
) -> AccountingFramework:
    """Crée une nouvelle version de référentiel ; une version existante n'est jamais modifiée."""
    _check(doc)
    pays = doc.get("pays") or "OHADA"
    exists = db.scalar(
        select(AccountingFramework).where(
            AccountingFramework.code == doc["code"],
            AccountingFramework.version == str(doc["version"]),
            AccountingFramework.pays == pays,
        )
    )
    if exists is not None:
        raise FrameworkError(
            f"La version {doc['version']} de {doc['code']} ({pays}) existe déjà : "
            "choisissez un nouveau numéro de version"
        )
    date_effet = doc["date_effet"]
    if isinstance(date_effet, str):
        date_effet = dt.date.fromisoformat(date_effet[:10])
    elif isinstance(date_effet, dt.datetime):
        date_effet = date_effet.date()
    fw = AccountingFramework(
        code=doc["code"],
        version=str(doc["version"]),
        libelle=doc["libelle"],
        pays=pays,
        date_effet=date_effet,
        source=doc.get("source"),
        fichier_sha256=sha256,
        ratios=doc.get("ratios", []),
        bases_seuil=doc.get("bases_seuil", []),
        regles=doc.get("regles", []),
        branches=doc.get("branches", []),
    )
    db.add(fw)
    db.flush()
    for acc in doc.get("comptes", []):
        db.add(
            FrameworkAccount(
                framework_id=fw.id,
                numero=str(acc["numero"]).strip(),
                libelle=acc["libelle"],
                classe=int(acc.get("classe") or str(acc["numero"])[0]),
                sens=acc.get("sens"),
                branche=acc.get("branche"),
            )
        )
    for i, p in enumerate(doc.get("postes", []), start=1):
        db.add(
            StatementLine(
                framework_id=fw.id,
                etat=p["etat"],
                code=p["code"],
                libelle=p["libelle"],
                nature=p["nature"],
                racines=p.get("racines", []),
                formule=p.get("formule") or None,
                ordre=int(p.get("ordre") or i),
                niveau=int(p.get("niveau") or 1),
            )
        )
    for param in doc.get("parametres", []):
        try:
            value = Decimal(str(param["valeur"]).replace(",", "."))
        except (InvalidOperation, KeyError) as exc:
            raise FrameworkError(
                f"Paramètre {param.get('code')} : valeur numérique attendue"
            ) from exc
        db.add(
            RegulatoryParameter(
                framework_id=fw.id,
                code=param["code"],
                libelle=param.get("libelle") or param["code"],
                valeur=value,
                unite=param.get("unite"),
                pays=param.get("pays"),
                exercice=int(param["exercice"]) if param.get("exercice") else None,
                branche=param.get("branche"),
            )
        )
    return fw


def builtin(name: str) -> dict[str, Any]:
    return json.loads((DATA_DIR / f"{name}.json").read_text(encoding="utf-8"))


# --- Excel ------------------------------------------------------------------------------------

SHEETS = {
    "Comptes": ("comptes", ["numero", "libelle", "classe", "sens", "branche"]),
    "Postes": (
        "postes",
        ["etat", "code", "libelle", "nature", "racines", "formule", "ordre", "niveau"],
    ),
    "Ratios": ("ratios", ["code", "libelle", "formule", "unite", "famille"]),
    "Bases": ("bases_seuil", ["code", "libelle", "formule", "pct_min", "pct_max"]),
    "Règles": ("regles", ["code", "libelle", "gauche", "droite"]),
    "Paramètres": (
        "parametres",
        ["code", "libelle", "valeur", "unite", "pays", "exercice", "branche"],
    ),
    "Branches": ("branches", ["code", "libelle"]),
}
HEADER_KEYS = ("code", "version", "libelle", "pays", "date_effet", "source")


def _racines_text(racines: list[dict[str, Any]]) -> str:
    return ";".join(r["racine"] + (f":{r['sens']}" if r.get("sens") else "") for r in racines)


def _racines_parse(text: Any) -> list[dict[str, Any]]:
    out = []
    for part in str(text or "").split(";"):
        part = part.strip()
        if not part:
            continue
        racine, _, sens = part.partition(":")
        out.append({"racine": racine.strip(), "sens": sens.strip() or None})
    return out


def read_excel(content: bytes) -> dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - fichier illisible
        raise FrameworkError("Classeur Excel illisible") from exc
    if "Référentiel" not in wb.sheetnames:
        raise FrameworkError("Feuille « Référentiel » absente")
    doc: dict[str, Any] = {}
    for row in wb["Référentiel"].iter_rows(values_only=True):
        if row and row[0] in HEADER_KEYS:
            doc[row[0]] = row[1]
    for sheet, (key, columns) in SHEETS.items():
        doc[key] = []
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        missing = [c for c in columns[:2] if c not in header]
        if missing:
            raise FrameworkError(f"Feuille « {sheet} » : colonne « {missing[0]} » absente")
        for row in rows[1:]:
            if not row or all(v is None for v in row):
                continue
            item = {c: row[header.index(c)] if c in header else None for c in columns}
            if key == "postes":
                item["racines"] = _racines_parse(item["racines"])
            if key == "comptes":
                item["numero"] = str(item["numero"]).split(".")[0]
            doc[key].append(item)
    return doc


def to_excel(db: Session, fw: AccountingFramework) -> bytes:
    wb = Workbook()
    head = wb.active
    head.title = "Référentiel"
    for key in HEADER_KEYS:
        head.append([key, str(getattr(fw, key)) if getattr(fw, key) is not None else None])
    accounts = db.scalars(
        select(FrameworkAccount)
        .where(FrameworkAccount.framework_id == fw.id)
        .order_by(FrameworkAccount.numero)
    )
    lines = db.scalars(
        select(StatementLine)
        .where(StatementLine.framework_id == fw.id)
        .order_by(StatementLine.ordre)
    )
    params = db.scalars(
        select(RegulatoryParameter).where(
            RegulatoryParameter.framework_id == fw.id, RegulatoryParameter.deleted_at.is_(None)
        )
    )
    data = {
        "comptes": [
            {
                "numero": a.numero,
                "libelle": a.libelle,
                "classe": a.classe,
                "sens": a.sens,
                "branche": a.branche,
            }
            for a in accounts
        ],
        "postes": [
            {
                "etat": p.etat,
                "code": p.code,
                "libelle": p.libelle,
                "nature": p.nature,
                "racines": _racines_text(p.racines),
                "formule": p.formule,
                "ordre": p.ordre,
                "niveau": p.niveau,
            }
            for p in lines
        ],
        "ratios": fw.ratios,
        "bases_seuil": fw.bases_seuil,
        "regles": fw.regles,
        "branches": fw.branches,
        "parametres": [
            {
                "code": p.code,
                "libelle": p.libelle,
                "valeur": float(p.valeur),
                "unite": p.unite,
                "pays": p.pays,
                "exercice": p.exercice,
                "branche": p.branche,
            }
            for p in params
        ],
    }
    for sheet, (key, columns) in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append(columns)
        for item in data[key]:
            ws.append([item.get(c) for c in columns])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
