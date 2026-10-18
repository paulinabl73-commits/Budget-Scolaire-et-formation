"""Circularisation (REV-19, ISA 505) : lettres de demande de confirmation en PDF et Word à
partir de modèles paramétrables (data/circularisation.json), et rapprochement automatique
de l'écart entre le solde comptable et le solde confirmé.

Le rapprochement cherche, parmi les lignes non lettrées du tiers dans le grand livre (REV-18),
une, deux ou trois lignes dont la somme explique exactement l'écart (facture ou règlement en
transit, par exemple) ; à défaut, l'écart reste à expliquer par l'auditeur.
"""

import io
import json
from functools import lru_cache
from itertools import combinations
from pathlib import Path
from typing import Any

from app.config import get_settings
from app.revision.third_parties import Item

MAX_COMBINED = 3
MAX_CANDIDATES = 40


@lru_cache
def templates() -> dict[str, Any]:
    path = Path(__file__).parent / "data" / "circularisation.json"
    return json.loads(path.read_text(encoding="utf-8"))


def _amount(value: Any) -> str:
    if value is None:
        return "……………"
    return f"{value:,.0f}".replace(",", " ")


def render(conf: Any, entity: str, closing: Any, currency: str) -> dict[str, Any]:
    s = get_settings()
    values = {
        "entite": entity,
        "date": closing.strftime("%d/%m/%Y"),
        "cabinet": s.nom_cabinet,
        "adresse_cabinet": s.adresse_cabinet or "",
        "tiers": conf.tiers,
        "compte": conf.compte or "",
        "solde": f"{_amount(conf.solde_comptable)} {currency}",
        "reference": f"CIRC-{conf.numero:03d}",
    }
    model = templates()[conf.type]
    return {
        "entite": entity,
        "destinataire": conf.tiers,
        "adresse": conf.adresse or "",
        "reference": values["reference"],
        "objet": model["objet"].format(**values),
        "paragraphes": [p.format(**values) for p in model["corps"]],
        "coupon": model["coupon"].format(**values),
    }


def to_pdf(letters: list[dict[str, Any]]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title="Circularisation",
    )
    story: list[Any] = []
    for i, letter in enumerate(letters):
        if i:
            story.append(PageBreak())
        story += [
            Paragraph(f"<b>{letter['entite']}</b>", styles["Normal"]),
            Spacer(1, 10 * mm),
            Paragraph(f"<b>{letter['destinataire']}</b>", styles["Normal"]),
        ]
        story += [Paragraph(line, styles["Normal"]) for line in letter["adresse"].splitlines()]
        story += [
            Spacer(1, 8 * mm),
            Paragraph(f"Référence : {letter['reference']}", styles["Normal"]),
            Paragraph(f"<b>Objet : {letter['objet']}</b>", styles["Normal"]),
            Spacer(1, 6 * mm),
        ]
        for paragraph in letter["paragraphes"]:
            story += [Paragraph(paragraph, styles["BodyText"]), Spacer(1, 3 * mm)]
        story += [
            Spacer(1, 12 * mm),
            Paragraph("La direction", styles["Normal"]),
            Spacer(1, 20 * mm),
            Paragraph("— — — — — — — — — — — — — — — — — — — —", styles["Normal"]),
            Paragraph(
                f"Réponse à retourner à nos auditeurs ({letter['reference']})", styles["Heading4"]
            ),
            Paragraph(letter["coupon"], styles["BodyText"]),
        ]
    doc.build(story)
    return out.getvalue()


def to_docx(letters: list[dict[str, Any]]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK

    document = Document()
    for i, letter in enumerate(letters):
        if i:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph().add_run(letter["entite"]).bold = True
        document.add_paragraph().add_run(letter["destinataire"]).bold = True
        for line in letter["adresse"].splitlines():
            document.add_paragraph(line)
        document.add_paragraph(f"Référence : {letter['reference']}")
        document.add_paragraph().add_run(f"Objet : {letter['objet']}").bold = True
        for paragraph in letter["paragraphes"]:
            document.add_paragraph(paragraph)
        document.add_paragraph("La direction")
        document.add_paragraph("—" * 30)
        document.add_heading(
            f"Réponse à retourner à nos auditeurs ({letter['reference']})", level=3
        )
        document.add_paragraph(letter["coupon"])
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def explain(gap: Any, items: list[Item], sign: int) -> list[Item] | None:
    """Lignes non lettrées dont la somme (vue de l'entité) égale l'écart ; les plus récentes
    sont essayées d'abord (opérations en transit à la clôture)."""
    candidates = sorted(items, key=lambda i: i.date or 0, reverse=True)[:MAX_CANDIDATES]
    for size in range(1, MAX_COMBINED + 1):
        for combo in combinations(candidates, size):
            if sum(sign * i.net for i in combo) == gap:
                return list(combo)
    return None
