"""Documents générés (note de synthèse, rapports, lettres) : une même description en blocs
est rendue en PDF (reportlab) ou en Word (python-docx).

Blocs : {"titre": str, "niveau": 1|2|3}, {"texte": str}, {"liste": [str]},
{"tableau": [[str]], "entete": bool}, {"saut": True} (saut de page), {"image": bytes,
"largeur_mm": float} (logo, REP-01), {"garde": str, "taille": int} (ligne centrée de la page de
garde). Un titre vide n'est pas répété en tête ; le pied de page est imprimé sur chaque page.
"""

import io
from typing import Any
from xml.sax.saxutils import escape


def to_pdf(title: str, blocks: list[dict[str, Any]], footer: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image,
        ListFlowable,
        ListItem,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
    )

    styles = getSampleStyleSheet()
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
    )
    story: list[Any] = [Paragraph(escape(title), styles["Title"])] if title else []
    for block in blocks:
        if "image" in block:
            img = Image(io.BytesIO(block["image"]))
            width = block.get("largeur_mm", 40) * mm
            img.drawHeight, img.drawWidth = width * img.imageHeight / img.imageWidth, width
            story += [img, Spacer(1, 8 * mm)]
        elif "garde" in block:
            size = block.get("taille", 12)
            style = ParagraphStyle(
                "garde", parent=styles["Normal"], alignment=TA_CENTER,
                fontSize=size, leading=size * 1.3, spaceAfter=size * 0.6,
            )  # fmt: skip
            story.append(Paragraph(escape(block["garde"]), style))
        elif "titre" in block:
            level = {1: "Heading1", 2: "Heading2"}.get(block.get("niveau", 2), "Heading3")
            story.append(Paragraph(escape(block["titre"]), styles[level]))
        elif "texte" in block:
            for line in str(block["texte"]).split("\n"):
                story.append(Paragraph(escape(line) or "&nbsp;", styles["BodyText"]))
            story.append(Spacer(1, 2 * mm))
        elif "liste" in block:
            items = [ListItem(Paragraph(escape(x), styles["BodyText"])) for x in block["liste"]]
            story.append(ListFlowable(items, bulletType="bullet"))
        elif "tableau" in block:
            rows = [
                [Paragraph(escape(str(c)), styles["BodyText"]) for c in r] for r in block["tableau"]
            ]
            if rows:
                table = Table(rows, repeatRows=1 if block.get("entete", True) else 0)
                style = [
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
                if block.get("entete", True):
                    style.append(("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke))
                table.setStyle(style)
                story += [table, Spacer(1, 3 * mm)]
        elif block.get("saut"):
            story.append(PageBreak())

    def page(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        if footer:
            canvas.drawString(18 * mm, 10 * mm, footer[:180])
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(document.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=page, onLaterPages=page)
    return out.getvalue()


def to_docx(title: str, blocks: list[dict[str, Any]], footer: str | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Mm, Pt

    document = Document()
    if footer:
        document.sections[0].footer.paragraphs[0].text = footer
    if title:
        document.add_heading(title, level=0)
    for block in blocks:
        if "image" in block:
            document.add_picture(io.BytesIO(block["image"]), width=Mm(block.get("largeur_mm", 40)))
        elif "garde" in block:
            para = document.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(block["garde"])
            run.font.size = Pt(block.get("taille", 12))
        elif "titre" in block:
            document.add_heading(block["titre"], level=block.get("niveau", 2))
        elif "texte" in block:
            for line in str(block["texte"]).split("\n"):
                document.add_paragraph(line)
        elif "liste" in block:
            for item in block["liste"]:
                document.add_paragraph(item, style="List Bullet")
        elif "tableau" in block and block["tableau"]:
            rows = block["tableau"]
            table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    table.cell(i, j).text = str(value)
                    if i == 0 and block.get("entete", True):
                        for run in table.cell(i, j).paragraphs[0].runs:
                            run.bold = True
        elif block.get("saut"):
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def render(fmt: str, title: str, blocks: list[dict[str, Any]], footer: str | None = None) -> bytes:
    return to_pdf(title, blocks, footer) if fmt == "pdf" else to_docx(title, blocks, footer)
