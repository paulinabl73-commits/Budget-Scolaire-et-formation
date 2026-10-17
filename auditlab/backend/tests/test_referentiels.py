"""REV-01 : référentiels versionnés, import et export Excel ; REV-03 : cadres normatifs."""

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.revision.formula import FormulaError, evaluate, references
from app.revision.frameworks import FrameworkError, builtin, load_framework
from app.security.roles import Role


def test_formules():
    values = {"AD": Decimal(100), "AJ": Decimal(50), "XB": Decimal(0)}
    assert evaluate("AD+AJ*2", values) == 200
    assert evaluate("(AD+AJ)/3", values) == 50
    assert evaluate("-AD+10", values) == -90
    assert evaluate("ZZ+1", values) == 1  # poste absent = 0
    assert evaluate("AD/XB", values) is None
    assert references("BI/XB*360") == {"BI", "XB"}
    for bad in ("AD+", "AD$1", "(AD", "AD AJ"):
        with pytest.raises(FormulaError):
            evaluate(bad, values)


def test_modele_syscohada_coherent():
    doc = builtin("syscohada_2017")
    codes = {p["code"] for p in doc["postes"]}
    assert {"AD", "BZ", "DZ", "XA", "XB", "XI", "CJ"} <= codes
    assert doc["ratios"] and doc["bases_seuil"] and doc["regles"]


def test_rev01_liste_detail_export_import(client, account, syscohada):
    admin = account(Role.ADMIN)
    auditor = account(Role.AUDITEUR)
    listed = client.get("/api/frameworks", headers=auditor.headers).json()
    assert [(f["code"], f["version"], f["nb_comptes"], f["nb_postes"]) for f in listed] == [
        ("SYSCOHADA", "2017", 160, 91)
    ]
    detail = client.get(f"/api/frameworks/{syscohada.id}", headers=auditor.headers).json()
    bi = next(p for p in detail["postes"] if p["code"] == "BI")
    assert {"racine": "41", "sens": "debit"} in bi["racines"]
    # Export, modification, réimport : nouvelle version, l'ancienne reste intacte.
    xlsx = client.get(f"/api/frameworks/{syscohada.id}/export", headers=auditor.headers)
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    head = wb["Référentiel"]
    for row in head.iter_rows():
        if row[0].value == "version":
            row[1].value = "2017-pays-a"
        if row[0].value == "pays":
            row[1].value = "Sénégal"
    wb["Comptes"].append(["4111", "Clients – ventes locales", 4, None, None])
    wb["Paramètres"].append(["taux_tva", "Taux normal de TVA", 0.18, "taux", "Sénégal", 2026, None])
    out = io.BytesIO()
    wb.save(out)
    files = {"file": ("ref.xlsx", out.getvalue())}
    assert (
        client.post("/api/frameworks/import", files=files, headers=auditor.headers).status_code
        == 403
    )
    resp = client.post("/api/frameworks/import", files=files, headers=admin.headers)
    assert resp.status_code == 201, resp.text
    new = resp.json()
    assert (new["version"], new["pays"], new["nb_comptes"]) == ("2017-pays-a", "Sénégal", 161)
    detail = client.get(f"/api/frameworks/{new['id']}", headers=admin.headers).json()
    assert detail["parametres"] == [{"code": "taux_tva", "libelle": "Taux normal de TVA",
                                     "valeur": "0.18000000", "unite": "taux", "pays": "Sénégal",
                                     "exercice": 2026, "branche": None}]  # fmt: skip
    racines = next(p for p in detail["postes"] if p["code"] == "BI")["racines"]
    assert {"racine": "41", "sens": "debit"} in racines  # « 41:debit » relu
    # La même version ne se réimporte pas.
    again = client.post("/api/frameworks/import", files=files, headers=admin.headers)
    assert again.status_code == 422 and "existe déjà" in again.text
    # Désactivation d'une version.
    resp = client.patch(
        f"/api/frameworks/{syscohada.id}", json={"actif": False}, headers=admin.headers
    )
    assert resp.json()["actif"] is False


@pytest.mark.parametrize(
    ("change", "message"),
    [
        (lambda d: d.update(code="XYZ"), "Code de référentiel"),
        (lambda d: d["comptes"].append({"numero": "4A", "libelle": "x", "classe": 4}), "invalide"),
        (lambda d: d["comptes"].append(dict(d["comptes"][0])), "en double"),
        (lambda d: d["postes"][0].update(formule="AD+ZZZ"), "poste inconnu"),
        (lambda d: d["postes"][0].update(nature="autre"), "nature invalide"),
        (
            lambda d: d["postes"][0]["racines"].append({"racine": "21", "sens": "gauche"}),
            "racine invalide",
        ),
        (
            lambda d: d["ratios"].append({"code": "x", "libelle": "x", "formule": "AD$"}),
            "non admis",
        ),
        (lambda d: d.update(parametres=[{"code": "t", "valeur": "abc"}]), "valeur numérique"),
    ],
)
def test_rev01_controles_du_chargement(db, change, message):
    doc = builtin("syscohada_2017")
    doc["version"] = "test"
    change(doc)
    with pytest.raises(FrameworkError, match=message):
        load_framework(db, doc)
    db.rollback()


def test_rev01_classeur_invalide(client, account):
    admin = account(Role.ADMIN)
    resp = client.post("/api/frameworks/import", files={"file": ("x.xlsx", b"pas un classeur")},
                       headers=admin.headers)  # fmt: skip
    assert resp.status_code == 422 and "illisible" in resp.text
    from openpyxl import Workbook

    wb = Workbook()
    out = io.BytesIO()
    wb.save(out)
    resp = client.post(
        "/api/frameworks/import", files={"file": ("x.xlsx", out.getvalue())}, headers=admin.headers
    )
    assert resp.status_code == 422 and "Référentiel" in resp.text
    assert (
        client.get(
            "/api/frameworks/00000000-0000-0000-0000-000000000000", headers=admin.headers
        ).status_code
        == 404
    )


def test_rev03_cadres_normatifs(client, account):
    admin = account(Role.ADMIN)
    chef = account(Role.CHEF_MISSION)
    resp = client.post("/api/normative-frameworks",
                       json={"code": "isa", "libelle": "Normes ISA"}, headers=admin.headers)  # fmt: skip
    assert resp.status_code == 201
    assert client.post("/api/normative-frameworks", json={"code": "isa", "libelle": "x"},
                       headers=admin.headers).status_code == 409  # fmt: skip
    assert client.post("/api/normative-frameworks", json={"code": "x", "libelle": "x"},
                       headers=chef.headers).status_code == 403  # fmt: skip
    nid = resp.json()["id"]
    resp = client.put(f"/api/normative-frameworks/{nid}",
                      json={"code": "isa", "libelle": "Normes ISA (IAASB)", "actif": True}, headers=admin.headers)  # fmt: skip
    assert resp.json()["libelle"] == "Normes ISA (IAASB)"
    assert [
        n["code"] for n in client.get("/api/normative-frameworks", headers=chef.headers).json()
    ] == ["isa"]
