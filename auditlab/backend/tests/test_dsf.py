"""REV-25 : aide à la liasse fiscale et statistique, modèles paramétrables par pays."""

import io

from openpyxl import Workbook, load_workbook

from app.models import DsfModel
from app.revision import dsf
from app.revision.frameworks import builtin
from app.security.roles import Role
from tests.revision_helpers import create_file, rev  # noqa: F401
from tests.test_affectation import load


def workbook(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tableau 9 – Résultat"
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_rev25_liasse(client, rev, db, account):  # noqa: F811
    generic = DsfModel(
        pays="OHADA",
        nom="Liasse OHADA – Système normal",
        version="2017",
        contenu=dsf.generic_model(builtin("syscohada_2017")["postes"]),
    )
    db.add(generic)
    db.commit()
    aud, admin = rev["aud"].headers, account(Role.ADMIN).headers
    models = client.get("/api/dsf-models", headers=aud).json()
    assert [f["nom"] for f in models[0]["feuilles"]] == [
        "Bilan – Actif",
        "Bilan – Passif",
        "Compte de résultat",
        "Tableau des flux de trésorerie",
    ]
    f = create_file(client, rev)
    url = f"/api/audit-files/{f['id']}/dsf?model_id={generic.id}"
    assert client.get(url, headers=aud).status_code == 409  # états impossibles sans balance
    load(client, rev, f["id"])
    sheets = {
        s["nom"]: {ln["ref"]: ln for ln in s["lignes"]}
        for s in client.get(url, headers=aud).json()["feuilles"]
    }
    assert (
        sheets["Bilan – Actif"]["BZ"]["N"] == "31700000.00"
        and sheets["Bilan – Actif"]["BZ"]["N-1"] == "23800000.00"
    )
    assert sheets["Tableau des flux de trésorerie"]["ZH"]["N-1"] is None
    resp = client.get(f"/api/audit-files/{f['id']}/dsf/export?model_id={generic.id}", headers=aud)
    wb = load_workbook(io.BytesIO(resp.content))
    assert (
        wb.sheetnames[0] == "Identification"
        and wb["Identification"]["B2"].value == "Société Ivoirienne de Négoce SA"
    )
    # Modèle d'un pays importé depuis Excel par l'administrateur.
    rows = [
        ["ref", "libelle", "note", "formule"],
        ["R1", "Résultat comptable", "N3", "XI"],
        ["R2", "Impôt sur le résultat", None, "RS"],
        ["R3", "Résultat avant impôt", None, "XI+RS"],
        ["R4", "Variation du résultat", None, "XI-XI_N1"],
    ]
    form = {"pays": "ci", "nom": "DSF Côte d'Ivoire – tableau 9", "version": "2026"}
    files = {"file": ("ci.xlsx", workbook(rows))}
    assert (
        client.post(
            "/api/dsf-models/import", data=form, files=files, headers=rev["chef"].headers
        ).status_code
        == 403
    )
    bad = {"file": ("bad.xlsx", workbook([*rows, ["R5", "Inconnu", None, "ZZ"]]))}
    resp = client.post("/api/dsf-models/import", data=form, files=bad, headers=admin)
    assert resp.status_code == 422 and "ZZ" in resp.text
    ci = client.post("/api/dsf-models/import", data=form, files=files, headers=admin).json()
    assert ci["pays"] == "CI" and ci["feuilles"] == [{"nom": "Tableau 9 – Résultat", "lignes": 4}]
    lines = client.get(f"/api/audit-files/{f['id']}/dsf?model_id={ci['id']}", headers=aud).json()[
        "feuilles"
    ][0]["lignes"]
    assert [(ln["ref"], ln["N"], ln["note"]) for ln in lines] == [
        ("R1", "7100000.00", "N3"),
        ("R2", "3000000.00", None),
        ("R3", "10100000.00", None),
        ("R4", "5600000.00", None),
    ]
    # Gabarit réimportable ; désactivation d'un modèle.
    template = client.get(f"/api/dsf-models/{ci['id']}/template", headers=aud)
    assert load_workbook(io.BytesIO(template.content)).sheetnames == ["Tableau 9 – Résultat"]
    assert (
        client.patch(f"/api/dsf-models/{ci['id']}?actif=false", headers=admin).json()["actif"]
        is False
    )
    assert (
        client.get(f"/api/audit-files/{f['id']}/dsf?model_id={f['id']}", headers=aud).status_code
        == 404
    )
