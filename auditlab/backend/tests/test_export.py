"""DAT-32 : export CSV et Excel."""

import datetime as dt
import io
from decimal import Decimal

import polars as pl
import pytest
from openpyxl import load_workbook

from app.analytics import export as export_module
from tests.data_helpers import make_dataset
from tests.test_datasets import ctx  # noqa: F401


@pytest.fixture
def table(db, ctx):  # noqa: F811
    df = pl.DataFrame(
        {
            "facture": ["F1", "F2", "F3"],
            "montant": pl.Series([Decimal("1250000.50"), Decimal("-45.25"), None]).cast(
                pl.Decimal(18, 2)
            ),
            "date": [dt.date(2026, 1, 31), dt.date(2026, 2, 1), None],
            "fournisseur": ["Bâtir & Co", "SOGECI; SARL", "Alpha"],
        }
    )
    return make_dataset(db, ctx["mission"]["id"], df, "Factures 2026 / test")


def test_dat32_csv_format_francais(client, ctx, table, db):  # noqa: F811
    from sqlalchemy import select

    from app.models import AuditLog

    resp = client.post(f"/api/datasets/{table.id}/export", json={"format": "csv"},
                       headers=ctx["aud"].headers)  # fmt: skip
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert 'filename="Factures_2026_test.csv"' in resp.headers["content-disposition"]
    raw = resp.content
    assert raw.startswith(b"\xef\xbb\xbf")
    lines = raw[3:].decode("utf-8").splitlines()
    assert lines[0] == "facture;montant;date;fournisseur"
    assert lines[1] == "F1;1250000,50;31/01/2026;Bâtir & Co"
    assert lines[2] == 'F2;-45,25;01/02/2026;"SOGECI; SARL"'
    assert lines[3] == "F3;;;Alpha"
    log = db.scalar(select(AuditLog).where(AuditLog.action == "dataset_exported"))
    assert log is not None and log.apres["lignes"] == 3


def test_dat32_excel_selection_et_informations(client, ctx, table):  # noqa: F811
    body = {
        "format": "xlsx",
        "colonnes": ["facture", "montant", "date"],
        "filtres": {"montant": {"filterType": "number", "type": "greaterThan", "filter": 0}},
        "tri": [{"colonne": "facture", "sens": "desc"}],
    }
    resp = client.post(f"/api/datasets/{table.id}/export", json=body, headers=ctx["aud"].headers)
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    data = wb["Données"]
    rows = list(data.iter_rows(values_only=True))
    assert rows[0] == ("facture", "montant", "date")
    assert rows[1][0] == "F1" and rows[1][1] == 1250000.5
    assert rows[1][2].date() == dt.date(2026, 1, 31)
    assert len(rows) == 2
    assert (
        data.cell(2, 2).number_format == "#,##0.00"
        and data.cell(2, 3).number_format == "DD/MM/YYYY"
    )
    info = dict(wb["Informations"].iter_rows(values_only=True))
    assert (
        info["Table"] == "Factures 2026 / test"
        and info["Empreinte SHA-256 de la table"] == table.parquet_sha256
    )
    assert info["Lignes exportées"] == 1 and info["Exporté par"]


def test_dat32_limites(client, ctx, table, monkeypatch):  # noqa: F811
    monkeypatch.setattr(export_module, "EXCEL_MAX_ROWS", 2)
    resp = client.post(
        f"/api/datasets/{table.id}/export", json={"format": "xlsx"}, headers=ctx["aud"].headers
    )
    assert resp.status_code == 422 and "exportez en CSV" in resp.text
    resp = client.post(f"/api/datasets/{table.id}/export", json={"format": "csv", "colonnes": ["x"]},
                       headers=ctx["aud"].headers)  # fmt: skip
    assert resp.status_code == 422
    resp = client.post(
        f"/api/datasets/{table.id}/export", json={"format": "pdf"}, headers=ctx["aud"].headers
    )
    assert resp.status_code == 422
