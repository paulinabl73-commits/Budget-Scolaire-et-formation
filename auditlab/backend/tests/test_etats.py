"""REV-23 et REV-24 : états financiers révisés, contrôles de cohérence, exports."""

import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import StatementLine
from tests.revision_helpers import create_file, rev  # noqa: F401
from tests.test_affectation import load
from tests.test_ajustements import PROVISION


def test_rev23_rev24_etats_revises(client, rev, db):  # noqa: F811
    f = create_file(client, rev)
    base = f"/api/audit-files/{f['id']}"
    aud = rev["aud"].headers
    assert client.get(f"{base}/statements", headers=aud).status_code == 409  # aucune balance
    load(client, rev, f["id"])
    adj = client.post(f"{base}/adjustments", json=PROVISION, headers=aud).json()
    client.post(f"/api/adjustments/{adj['id']}/decision", json={"statut": "acceptee"}, headers=aud)
    data = client.get(f"{base}/statements", headers=aud).json()
    assert list(data["etats"]) == [
        "bilan_actif",
        "bilan_passif",
        "compte_resultat",
        "flux_tresorerie",
    ]
    actif = {x["code"]: x for x in data["etats"]["bilan_actif"]}
    assert (actif["BI"]["N_avant"], actif["BI"]["ajustements"], actif["BI"]["N"]) == (
        "7000000.00", "-500000.00", "6500000.00")  # fmt: skip
    resultat = {x["code"]: x for x in data["etats"]["compte_resultat"]}
    assert resultat["XI"]["N_avant"] == "7100000.00" and resultat["XI"]["N"] == "6600000.00"
    flux = {x["code"]: x for x in data["etats"]["flux_tresorerie"]}
    assert flux["ZA"]["N"] == "6800000.00" and flux["ZT"]["N"] == "13700000.00"
    assert flux["ZH"]["N"] == flux["ZT"]["N"] and flux["ZG"]["N"] == "6900000.00"
    assert flux["ZA"]["N-1"] is None  # le tableau N-1 supposerait le bilan N-2
    assert [c["code"] for c in data["controles"]] == ["bilan_equilibre", "resultat_concordant",
                                                      "tresorerie_concordante"]  # fmt: skip
    assert all(c["ok"] for c in data["controles"]) and list(data["controles"][2]["ecarts"]) == ["N"]
    # Exports Excel et PDF.
    resp = client.get(f"{base}/statements/export?format=xlsx", headers=aud)
    assert (
        resp.status_code == 200
        and "etats_financiers_2026.xlsx" in resp.headers["content-disposition"]
    )
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Bilan – actif", "Bilan – passif", "Compte de résultat",
                             "Tableau des flux de trésorerie", "Contrôles"]  # fmt: skip
    rows = {r[0]: r for r in wb["Bilan – actif"].iter_rows(min_row=5, values_only=True)}
    assert rows["BI"][2:5] == (7000000, -500000, 6500000)
    pdf = client.get(f"{base}/statements/export?format=pdf", headers=aud)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert client.get(f"{base}/statements/export?format=doc", headers=aud).status_code == 422
    # Un état incohérent (total de l'actif immobilisé amputé d'un poste) n'est pas exporté.
    total = db.scalar(select(StatementLine).where(StatementLine.code == "AZ"))
    total.formule = total.formule.replace("+AN", "")
    db.commit()
    data = client.get(f"{base}/statements", headers=aud).json()
    assert not data["controles"][0]["ok"]
    resp = client.get(f"{base}/statements/export?format=pdf", headers=aud)
    assert resp.status_code == 409 and "bilan est équilibré" in resp.text
