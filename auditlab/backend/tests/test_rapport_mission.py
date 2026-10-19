"""REP-02 : génération du rapport de mission en PDF et Word."""

import io

import pytest
from docx import Document

from app.reporting import templates
from tests.finding_helpers import create_finding, team  # noqa: F401
from tests.report_helpers import validated_finding


@pytest.fixture
def modele(db):
    t = templates.seed(db)
    db.commit()
    return t


def _text(content):
    doc = Document(io.BytesIO(content))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_rep02_rapport(client, team, modele):  # noqa: F811
    chef, aud = team["chef"].headers, team["aud"].headers
    validated_finding(client, team, "moyenne", "Retards de facturation")
    validated_finding(client, team, "critique", "Paiements sans pièce")
    create_finding(client, team, titre="Constat encore en brouillon")
    url = f"{team['url']}/report"
    report = client.get(url, headers=aud).json()
    assert report["modele"] == modele.nom and report["peut_rediger"]
    body = {
        "synthese": "Le contrôle interne des achats est insuffisant.",
        "conclusion": "Plan à suivre.",
    }
    assert client.put(url, json=body, headers=team["da"].headers).status_code == 403
    assert client.put(url, json=body, headers=chef).json()["synthese"].startswith("Le contrôle")
    docx = client.get(f"{url}/document?format=docx", headers=team["directeur"].headers)
    assert docx.status_code == 200
    text = _text(docx.content)
    for expected in ("Rapport d'audit interne", "Diffusion restreinte", "Le contrôle interne",
                     "Criticité : Critique", "Accord partiel : 8 créations", "Réactiver le blocage",
                     "PROJET – constats en phase contradictoire non repris", "Papiers de travail"):  # fmt: skip
        assert expected in text, expected
    assert "Constat encore en brouillon" not in text.split("PROJET")[0]
    # Ordre par criticité décroissante : le constat critique précède le constat moyen.
    assert text.index("Paiements sans pièce") < text.index("Retards de facturation")
    pdf = client.get(f"{url}/document", headers=chef)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"


def test_rep02_logo(client, team, modele):  # noqa: F811
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (120, 40), (31, 95, 168)).save(buf, format="PNG")
    client.put(f"/api/mission-report-templates/{modele.id}/logo",
               files={"file": ("logo.png", io.BytesIO(buf.getvalue()))}, headers=team["directeur"].headers)  # fmt: skip
    for fmt in ("pdf", "docx"):
        resp = client.get(
            f"{team['url']}/report/document?format={fmt}", headers=team["chef"].headers
        )
        assert resp.status_code == 200, fmt
    doc = Document(io.BytesIO(resp.content))
    assert len(doc.inline_shapes) == 1


def test_rep03_circuit(client, team, modele, account, db):  # noqa: F811
    from app.models import Engagement

    chef, aud, directeur = team["chef"].headers, team["aud"].headers, team["directeur"].headers
    validated_finding(client, team, "critique", "Paiements sans pièce")
    url = f"{team['url']}/report"
    report = client.get(url, headers=chef).json()
    assert report["statut"] == "projet" and report["transitions"] == []  # le chef ne se revoit pas
    move = f"{url}/transition"
    assert client.post(move, json={"statut": "revu"}, headers=chef).status_code == 403
    revu = client.post(move, json={"statut": "revu"}, headers=directeur).json()
    assert revu["statut"] == "revu" and not client.get(url, headers=chef).json()["peut_rediger"]
    assert client.put(url, json={"synthese": "x"}, headers=chef).status_code == 409
    assert client.post(move, json={"statut": "projet"}, headers=directeur).status_code == 422
    back = client.post(move, json={"statut": "projet", "commentaire": "Préciser la synthèse"},
                       headers=directeur).json()  # fmt: skip
    assert (
        back["statut"] == "projet"
        and back["historique"][-1]["commentaire"] == "Préciser la synthèse"
    )
    client.put(url, json={"synthese": "Synthèse précisée."}, headers=chef)
    client.post(move, json={"statut": "revu"}, headers=directeur)
    assert client.post(move, json={"statut": "valide"}, headers=aud).status_code == 403
    assert (
        client.post(move, json={"statut": "valide"}, headers=directeur).json()["statut"] == "valide"
    )
    # Diffusion : mission au statut « rapport » (constats validés).
    assert client.post(move, json={"statut": "diffuse"}, headers=chef).status_code == 409
    db.get(Engagement, team["mission"]["id"]).statut = "rapport"
    db.commit()
    diffused = client.post(move, json={"statut": "diffuse"}, headers=chef).json()
    assert (
        diffused["statut"] == "diffuse" and diffused["diffuse_le"] and diffused["transitions"] == []
    )
    assert set(diffused["empreintes"]) == {"pdf", "docx"}
    first = client.get(f"{url}/document", headers=chef).content
    # Version figée : le texte de la mission change, le rapport diffusé non.
    import hashlib

    assert hashlib.sha256(first).hexdigest() == diffused["empreintes"]["pdf"]
    modele.page_garde = {**modele.page_garde, "titre": "Autre titre"}
    db.commit()
    assert client.get(f"{url}/document", headers=chef).content == first
    assert client.put(url, json={"synthese": "y"}, headers=chef).status_code == 409
    # Le rapport n'empêche plus la clôture (reste le papier A-1 du jeu de test à approuver).
    mission = client.get(team["url"], headers=chef).json()
    blocked = next(t for t in mission["transitions"] if t["statut"] == "cloturee")
    assert blocked["motif_blocage"] == "Papiers de travail non approuvés : A-1"


def test_rep03_cloture_sans_rapport(client, team, modele, db):  # noqa: F811
    from app.models import Engagement

    db.get(Engagement, team["mission"]["id"]).statut = "rapport"
    db.commit()
    mission = client.get(team["url"], headers=team["chef"].headers).json()
    blocked = next(t for t in mission["transitions"] if t["statut"] == "cloturee")
    assert blocked["motif_blocage"] == "Le rapport de mission doit être diffusé avant la clôture"
