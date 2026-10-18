"""REV-19 : circularisation (sélection, lettres PDF et Word, suivi, rapprochement des écarts)."""

import io

from docx import Document

from tests.revision_helpers import create_file, rev  # noqa: F401
from tests.test_balances import upload
from tests.test_tiers import ledger_and_balance


def test_rev19_circularisation(client, rev):  # noqa: F811
    f = create_file(client, rev)
    base = f"/api/audit-files/{f['id']}"
    aud = rev["aud"].headers
    gl, bal = ledger_and_balance()
    upload(client, rev, f["id"], "bal.csv", bal)
    upload(client, rev, f["id"], "gl.csv", gl, type="grand_livre")
    # Sélection automatique par tiers (soldes vus de l'entité), puis banques et avocat.
    assert (
        client.post(
            f"{base}/confirmations/select", json={"type": "client"}, headers=rev["da"].headers
        ).status_code
        == 403
    )
    sel = client.post(
        f"{base}/confirmations/select", json={"type": "client", "seuil": "0"}, headers=aud
    ).json()
    assert sel["ajoutees"] == 3
    clients = {c["tiers"]: c for c in sel["confirmations"]}
    assert {k: v["solde_comptable"] for k, v in clients.items()} == {
        "A": "590000.00",
        "B": "136000.00",
        "C": "18000.00",
    }
    assert [c["reference"] for c in sel["confirmations"]] == ["CIRC-001", "CIRC-002", "CIRC-003"]
    again = client.post(
        f"{base}/confirmations/select", json={"type": "client", "seuil": "0"}, headers=aud
    ).json()
    assert again["ajoutees"] == 0  # pas de doublon
    sup = client.post(
        f"{base}/confirmations/select", json={"type": "fournisseur", "seuil": "0"}, headers=aud
    ).json()
    supplier = next(c for c in sup["confirmations"] if c["type"] == "fournisseur")
    assert supplier["tiers"] == "S" and supplier["solde_comptable"] == "200000.00"
    banks = client.post(f"{base}/confirmations/select", json={"type": "banque"}, headers=aud).json()
    assert banks["ajoutees"] == 1
    lawyer = client.post(f"{base}/confirmations", headers=aud, json={
        "type": "avocat", "tiers": "Maître Konan", "adresse": "Plateau, Abidjan"}).json()  # fmt: skip
    assert lawyer["reference"] == "CIRC-006"
    # Lettres PDF (toutes) et Word (une demande).
    pdf = client.get(f"{base}/confirmations/letters?format=pdf", headers=aud)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert "circularisation_2026.pdf" in pdf.headers["content-disposition"]
    word = client.get(f"/api/confirmations/{clients['A']['id']}/letter?format=docx", headers=aud)
    text = "\n".join(p.text for p in Document(io.BytesIO(word.content)).paragraphs)
    assert "CIRC-001" in text and "590 000 XOF en notre faveur" in text and "31/12/2026" in text
    lawyer_text = client.get(f"/api/confirmations/{lawyer['id']}/letter?format=docx", headers=aud)
    assert "litiges" in "\n".join(
        p.text for p in Document(io.BytesIO(lawyer_text.content)).paragraphs
    )
    # Suivi : envoi, réponses, rapprochement automatique des écarts.
    for c in sel["confirmations"]:
        sent = client.patch(
            f"/api/confirmations/{c['id']}", json={"statut": "envoyee"}, headers=aud
        ).json()
        assert sent["envoyee_le"] is not None
    assert (
        client.get(f"{base}/confirmations/letters?format=pdf&type=client", headers=aud).status_code
        == 409
    )
    a = client.patch(
        f"/api/confirmations/{clients['A']['id']}", json={"solde_confirme": "0"}, headers=aud
    ).json()
    assert a["statut"] == "recue" and a["ecart"] == "590000.00"
    rec = client.post(f"/api/confirmations/{clients['A']['id']}/reconcile", headers=aud).json()
    assert [ln["piece"] for ln in rec["lignes"]] == ["F2"] and rec["explication"].startswith(
        "Écart expliqué"
    )
    client.patch(
        f"/api/confirmations/{clients['B']['id']}", json={"solde_confirme": "236000"}, headers=aud
    )
    rec = client.post(f"/api/confirmations/{clients['B']['id']}/reconcile", headers=aud).json()
    assert [ln["piece"] for ln in rec["lignes"]] == ["R3"]  # règlement en transit
    client.patch(
        f"/api/confirmations/{clients['C']['id']}", json={"solde_confirme": "17000"}, headers=aud
    )
    rec = client.post(f"/api/confirmations/{clients['C']['id']}/reconcile", headers=aud).json()
    assert rec["lignes"] == [] and "non expliqué" in rec["message"]
    assert (
        client.post(f"/api/confirmations/{supplier['id']}/reconcile", headers=aud).status_code
        == 409
    )
    synthese = client.get(f"{base}/confirmations", headers=aud).json()["synthese"]["client"]
    assert synthese == {"envoyees": 3, "recues": 3, "taux_reponse": 100.0, "montant_envoye": "744000.00",
                        "montant_recu": "744000.00", "ecarts": 3, "ecarts_non_expliques": 1}  # fmt: skip
    # Suppression : seulement avant l'envoi.
    assert client.delete(f"/api/confirmations/{clients['A']['id']}", headers=aud).status_code == 409
    assert client.delete(f"/api/confirmations/{lawyer['id']}", headers=aud).status_code == 204
